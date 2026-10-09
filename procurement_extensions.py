from pathlib import Path
from io import BytesIO
from datetime import datetime, timezone
import re
import pandas as pd
from openpyxl import load_workbook

TEMPLATE = Path(__file__).resolve().parent/'templates'/'AIRB-PD-F02- Purchase Requisition (PR) Form rev.5.xlsx'
STAGES=['New','Pending Engineer Review','Pending Approval','Submitted to Procurement','RFQ in Progress','Quotation Received','PO Issued','Awaiting Delivery','Partially Delivered','Delivered / Completed','Closed','Rejected / Cancelled']
WRITERS={'Administrator','Procurement'}
READERS=WRITERS|{'Engineer','Business Development'}
def now():return datetime.now(timezone.utc).isoformat()
def generate_pr_xlsx(data):
    wb=load_workbook(TEMPLATE);ws=wb['REV.5']
    fields={'T3':'pr_no','T4':'project','D5':'requested_by','T5':'date_required','D6':'site','L6':'request_date','B16':'justification'}
    def put(addr,value):
        cell=ws[addr]
        if cell.__class__.__name__=='MergedCell':
            for merged in ws.merged_cells.ranges:
                if addr in merged:cell=ws.cell(merged.min_row,merged.min_col);break
        cell.value=value
    for addr,key in fields.items():
        if data.get(key):put(addr,data[key])
    items=data.get('items',[])
    if len(items)>10:raise ValueError('The AIRB form supports 10 item rows.')
    for i,item in enumerate(items):
        r=21+i
        for col,key in [('B','account_code'),('D','description'),('N','unit'),('O','quantity'),('Q','unit_price')]:
            put(f'{col}{r}',item.get(key,''))
        put(f'T{r}',f'=O{r}*Q{r}')
    put('T31','=SUM(T21:T30)')
    stream=BytesIO();wb.save(stream);return stream.getvalue()
def render_pr_documents(st,db,requests,user):
    st.subheader('Generate AIRB PR (Rev.05)')
    if not requests:
        st.info('Save a procurement request first.');return
    ids=[str(r.get('Request ID')) for r in requests if r.get('Request ID')]
    if not ids:return
    selected=st.selectbox('Request',ids,key='pr_req');req=next(r for r in requests if str(r.get('Request ID'))==selected)
    with st.form('pr_form'):
        c1,c2=st.columns(2)
        with c1:
            pr_no=st.text_input('PR number',value=str(req.get('PR No.',selected)))
            project=st.text_input('Project name')
            requested_by=st.text_input('Requested by',value=str(req.get('Requested By',user.get('full_name',''))))
            site=st.text_input('Department / site',value=str(req.get('Site ID','')))
        with c2:
            request_date=st.date_input('Request date');date_required=st.date_input('Date required')
            quantity=st.number_input('Quantity',min_value=0.01,value=1.0)
            price=st.number_input('Estimated unit price RM',min_value=0.0,value=float(req.get('Estimated Cost') or 0))
        description=st.text_area('Item / service description',value=str(req.get('Requirement','')))
        justification=st.text_area('Reason for purchase',value=str(req.get('Justification','')))
        unit=st.text_input('Unit',value='unit')
        generate=st.form_submit_button('Prepare AIRB PR Excel')
    if generate:
        try:
            st.session_state['pr_download']=generate_pr_xlsx({'pr_no':pr_no,'project':project,'requested_by':requested_by,'site':site,'request_date':request_date.strftime('%d/%m/%Y'),'date_required':date_required.strftime('%d/%m/%Y'),'justification':justification,'items':[{'description':description,'unit':unit,'quantity':quantity,'unit_price':price}]})
            st.session_state['pr_filename']='AIRB_PR_'+re.sub('[^A-Za-z0-9_-]','_',pr_no)+'.xlsx'
        except Exception as e:st.error(f'PR generation failed: {e}')
    if st.session_state.get('pr_download'):
        st.download_button('Download AIRB PR',st.session_state['pr_download'],file_name=st.session_state['pr_filename'])
    st.caption('Draft only. Finance and LOA approvals/signatures are not generated.')
    st.subheader('Procurement handover / progress')
    current=req.get('Status','New');st.write('Current status:',current)
    if user.get('role') in {'Administrator','Engineer'} and st.button('Submit to Procurement'):
        if current in {'New','Pending Engineer Review','Pending Approval','PR / IER Issued'}:
            updated={**req,'Status':'Submitted to Procurement','Submitted By':user.get('full_name',''),'Submitted At':now()}
            try:
                result=db.table('procurement_requests').update({'request_data':updated}).eq('request_id',selected).eq('site_id',req.get('Site ID')).execute()
                if not result.data:raise ValueError('No matching request')
                req.update(updated);st.success('Submitted. No email notification is configured.')
            except Exception as e:st.error(str(e))
        else:st.warning('Request is already in procurement processing.')
    if user.get('role') in WRITERS:
        stage=st.selectbox('Procurement stage',STAGES,index=STAGES.index(current) if current in STAGES else 0)
        supplier=st.text_input('Supplier',value=str(req.get('Supplier','')))
        po=st.text_input('PO number',value=str(req.get('PO No.','')))
        remarks=st.text_area('Progress remarks')
        if st.button('Save procurement status'):
            history=list(req.get('Status History',[]));history.append({'stage':stage,'by':user.get('full_name',''),'at':now(),'remarks':remarks})
            updated={**req,'Status':stage,'Supplier':supplier,'PO No.':po,'Status History':history}
            try:
                result=db.table('procurement_requests').update({'request_data':updated}).eq('request_id',selected).eq('site_id',req.get('Site ID')).execute()
                if not result.data:raise ValueError('No matching request')
                req.update(updated);st.success('Saved')
            except Exception as e:st.error(str(e))
def render_price_library(st,db,user):
    st.title('Procurement Price Library')
    if user.get('role') not in READERS:st.error('Access denied');return
    try:rows=db.table('procurement_price_history').select('*').order('purchased_on',desc=True).limit(2000).execute().data or []
    except Exception as e:st.error(f'Run SUPABASE_SETUP.sql first: {e}');return
    search=st.text_input('Search item, specification, supplier, PO')
    matches=[r for r in rows if search.lower() in ' '.join(str(r.get(k,'')) for k in ('item_name','specification','supplier','po_no','model')).lower()]
    if matches:
        df=pd.DataFrame(matches);cols=[x for x in ['purchased_on','item_name','specification','brand','model','supplier','quantity','unit','unit_price','currency','po_no','site_id'] if x in df]
        st.dataframe(df[cols],use_container_width=True,hide_index=True)
        st.download_button('Export price history CSV',df[cols].to_csv(index=False),file_name='price_history.csv')
        prices=[float(r['unit_price']) for r in matches if r.get('unit_price') is not None and r.get('currency')=='MYR']
        if prices:
            a,b,c=st.columns(3);a.metric('Lowest MYR price',f'RM {min(prices):,.2f}');b.metric('Average MYR price',f'RM {sum(prices)/len(prices):,.2f}');c.metric('Latest MYR price',f'RM {prices[0]:,.2f}')
        st.caption('Compare equivalent specifications and purchase quantities before budgeting.')
    else:st.info('No purchases found.')
    if user.get('role') not in WRITERS:st.info('Read-only: Procurement / Administrator maintain prices.');return
    st.subheader('Add confirmed purchase')
    with st.form('purchase_entry'):
        a,b=st.columns(2)
        with a:
            item=st.text_input('Item name *');spec=st.text_input('Specification');brand=st.text_input('Brand');model=st.text_input('Model');supplier=st.text_input('Supplier *');po=st.text_input('PO No. *')
        with b:
            date=st.date_input('Purchase date');qty=st.number_input('Quantity',min_value=0.01,value=1.0);unit=st.text_input('Unit',value='unit');price=st.number_input('Actual unit price',min_value=0.0);currency=st.selectbox('Currency',['MYR','USD','EUR','SGD','GBP','CNY']);site=st.text_input('Site ID');pr=st.text_input('PR No.')
        save=st.form_submit_button('Record purchase')
    if save:
        if not item.strip() or not supplier.strip() or not po.strip():st.error('Item, supplier and PO required');return
        data={'item_name':item,'specification':spec,'brand':brand,'model':model,'supplier':supplier,'po_no':po,'pr_no':pr,'purchased_on':date.isoformat(),'quantity':qty,'unit':unit,'unit_price':price,'currency':currency,'site_id':site,'recorded_by':user.get('full_name',''),'recorded_at':now()}
        try:db.table('procurement_price_history').insert(data).execute();st.success('Recorded');st.rerun()
        except Exception as e:st.error(f'Unable to record purchase: {e}')
    st.caption('Historical entries are append-only in this interface.')


def render_linked_po_workflow(st, db, requests, user):
    """PO lifecycle is attached to an existing saved PR/request, never standalone."""
    st.subheader('Linked Purchase Orders')
    st.caption('Select an existing procurement request. A PO inherits its PR, site, work order and IER references.')
    eligible = [r for r in requests if r.get('Request ID')]
    if not eligible:
        st.info('Create and save a PR request before recording a PO.')
        return
    ids = [str(r['Request ID']) for r in eligible]
    chosen = st.selectbox('Linked PR / procurement request', ids, key='linked_po_req')
    req = next(r for r in eligible if str(r['Request ID']) == chosen)
    st.write('PR:', req.get('PR No.') or chosen, '| Site:', req.get('Site ID') or 'Not set')
    st.write('WO:', req.get('WO ID') or '—', '| CM:', req.get('CM ID') or '—', '| IER:', req.get('IER Report No.') or '—')
    orders = list(req.get('Purchase Orders') or [])
    if orders:
        st.dataframe(pd.DataFrame(orders), hide_index=True, use_container_width=True)
    if user.get('role') not in WRITERS:
        st.info('Read-only: Procurement and Administrator maintain purchase orders.')
        return
    site_id = str(req.get('Site ID') or '').strip()
    if not site_id or site_id == 'ALL':
        st.error('This request needs a valid operational Site ID before PO processing.')
        return
    with st.form('linked_po_form'):
        po_no = st.text_input('PO number *')
        supplier = st.text_input('Supplier *')
        description = st.text_input('Item / service description *', value=str(req.get('Requirement') or ''))
        specification = st.text_input('Specification / model')
        quantity = st.number_input('PO quantity', min_value=0.01, value=1.0)
        unit = st.text_input('Unit', value='unit')
        price = st.number_input('Agreed unit price (RM)', min_value=0.0, value=0.0)
        stage = st.selectbox('PO status', ['PO Issued','Awaiting Delivery','Partially Delivered','Delivered / Completed'])
        received = st.number_input('Quantity received to date', min_value=0.0, max_value=float(quantity), value=0.0)
        note = st.text_area('Remarks / delivery reference')
        submit = st.form_submit_button('Save linked PO')
    if submit:
        if not all(x.strip() for x in (po_no, supplier, description)):
            st.error('PO number, supplier and description are required.')
            return
        if any(str(o.get('po_no')) == po_no and str(o.get('description')) == description for o in orders):
            st.error('This PO line already exists. Update the existing line instead of duplicating it.')
            return
        if stage == 'Delivered / Completed' and received < quantity:
            st.error('Received quantity must equal ordered quantity before marking fully delivered.')
            return
        order = {'po_no':po_no.strip(),'supplier':supplier.strip(),'description':description.strip(),
                 'specification':specification.strip(),'quantity':float(quantity),'unit':unit,
                 'unit_price':float(price),'currency':'MYR','status':stage,
                 'received_quantity':float(received),'remarks':note,'recorded_at':now(),
                 'recorded_by':user.get('full_name','')}
        updated = {**req,'Purchase Orders':orders+[order], 'Status':stage,
                   'Status History':list(req.get('Status History') or []) +
                   [{'stage':stage,'by':user.get('full_name',''),'at':now(),'remarks':f'PO {po_no}: {note}'}]}
        try:
            result = db.table('procurement_requests').update({'request_data':updated}).eq('request_id',chosen).eq('site_id',site_id).execute()
            if not result.data:
                raise ValueError('No matching PR in the database. Nothing was saved.')
            req.update(updated)
            st.success('PO linked to the original procurement request.')
            st.rerun()
        except Exception as exc:
            st.error(f'Unable to save linked PO: {exc}')
    st.caption('Price Library entry is confirmed separately upon delivery/invoice verification; a PO alone is not proof of purchase.')
